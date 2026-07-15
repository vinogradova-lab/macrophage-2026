""" Use conda env "polars"

Collects a raw p-value and log2 fold change for target proteins across four modalities
(protein/TMT, RNA-seq, western blot, qPCR) into one tidy CSV. Source files are read-only.

Protein (TMT volcano table) and RNA (DESeq2, raw Wald pvalue) are copied through
pre-computed. Western blot and qPCR are raw replicate values (3 reps, D1/D2/D3) that we
test here, both comparing treatment vs the M0 control:
- western_blot: reps are pre-normalized to M0=1, so log2 ratios get a ONE-SAMPLE t-test
  vs 0 (ttest_1samp) -- equivalent to a paired test since M0 is a constant 0 baseline.
- qpcr: log2 values with a per-replicate M0, so a PAIRED t-test vs M0 (ttest_rel).
A condition with fewer than 2 usable replicates is dropped to a null p-value.

Outputs: OUTPUT_CSV, a synced copy in MM_HEATMAP_DIR for the R heatmap, and
TBCK_OUTPUT_CSV (adds the reactivity modality for the TBCK complex).
"""

import math
import re

import openpyxl
import polars as pl
from pathlib import Path
from scipy import stats

from src.macrophage import (
    CONDITIONS,
    DEG_FILE,
    DEG_SHEET,
    FIGURE1_PANELS,
    RNA_DEG_DIR,
)

# Figure 1 "Panels/" holds the cross-modal source tables (volcano long-format, RNA DEG,
# western blot / qPCR source data); Figure 5 is a sibling two levels up.
data_dir = FIGURE1_PANELS

PROTEINS_TO_CALC = ["CAMK1", "SCO2", "STAT4", "S100A4"]
# TBCK complex — only RNA-seq and whole proteome exist for these (no WB/qPCR).
TBCK_COMPLEX_PROTEINS = ["TBCK", "CRYZL1", "C12orf4", "PPP1R21", "GATD1"]
TBCK_OUTPUT_CSV = (data_dir.parent.parent
                   / "Figure 5/Panels/multi_modal_heatmap/multi_modal_significance.csv")

# Protein and RNA already has P-values, FC calculated (T-test and DESeq)
PROTEIN_FC_P_VALUES = data_dir / 'volcano_plot/2026_ontology/volcano_data_long_format.csv'
# separate excel files for each condition
RNA_FC_P_VALUES_FOLDER = RNA_DEG_DIR

# Reactivity — residue-level cysteine reactivity (percent of M0 control). Significance is
# a categorical per-residue call (direction_of_reactivity_change), not a p-value.
REACTIVITY_FC_P_VALUES = Path(
    "/Users/henrysanford/Dropbox @RU Dropbox/Vinogradova Laboratory/Macrophage project/"
    "Manuscript 1/01_Data Analysis/01_Proteomics/02_reactivity/03_rc_analysis/"
    "reactivity_output/20251023_ratio_directionality/median_for_5+_peptides/output/"
    "reactivity_changes_long_format.csv")

# WB data is normalized several different ways: for CAMK1/SCO2/STAT4/CLIC1 the final
# value is "D* Vinc/M0 norm", for S100A4 it is "double norm" (see WB_CONFIG below).
WESTERN_BLOT_REPLICATE_VALUES = data_dir / "20260709 TLR-stim compiled source data.xlsx"
# qpcr data under qPCR sheet_name (moved into the Figure S2B-D subfolder)
QPCR_REPLICATE_VALUES = data_dir / "Figure S2B-D" / "protein vs rna examples.xlsx"

OUTPUT_CSV = Path(__file__).resolve().parent / "multi_modal_significance.csv"
# The R heatmap chunk (wp_visualization.Rmd) reads the CSV from this dir, so write a
# copy there too, keeping the local (git-tracked) copy and the figure copy in sync.
MM_HEATMAP_DIR = data_dir / "multi_modal_heatmap"

# The seven treatment conditions shared across modalities (canonical order in src.macrophage).
SHARED_CONDITIONS = CONDITIONS

# Final schema for every per-modality DataFrame. Fold change is log2 for every modality
# (hence the column name); p_value is the raw test p-value, not a -log10 transform.
SCHEMA = ["protein", "condition", "modality", "p_value", "log2_fold_change"]


def _norm_cond(label) -> str | None:
    """Map the various condition spellings to the shared vocabulary."""
    if label is None:
        return None
    label = str(label).strip()
    # Western blot uses "TLR1/2" (S100A4 sheet) or "TLR1" (others); RNA file uses TLR1.
    return "TLR1-2" if label in ("TLR1", "TLR1/2", "TLR1-2", "TLR1_2") else label


def _row(protein, condition, modality, p, fc) -> dict:
    p = None if p is None or (isinstance(p, float) and math.isnan(p)) else float(p)
    fc = None if fc is None else float(fc)
    return {"protein": protein, "condition": condition, "modality": modality,
            "p_value": p, "log2_fold_change": fc}


def _wb(name) -> openpyxl.Workbook:
    return openpyxl.load_workbook(name, read_only=True, data_only=True)


# Protein (TMT) — already has p_value and log2_FC, just reshape/filter.
# Match protein names case-insensitively and emit the canonical requested name so the
# same protein lines up with the RNA modality (e.g. C12orf4 vs C12ORF4).
def load_protein(proteins=PROTEINS_TO_CALC) -> pl.DataFrame:
    canonical = {p.upper(): p for p in proteins}
    return (
        pl.read_csv(PROTEIN_FC_P_VALUES, infer_schema_length=2000)
        .filter(pl.col("protein").str.to_uppercase().is_in(list(canonical))
                & pl.col("condition").is_in(SHARED_CONDITIONS))
        .select(
            pl.col("protein").str.to_uppercase().replace(canonical).alias("protein"),
            "condition",
            pl.lit("protein").alias("modality"),
            pl.col("p_value"),
            pl.col("log2_FC").alias("log2_fold_change"),
        )
        .select(SCHEMA)
    )


# RNA-seq — one DESeq2 xlsx per condition (src.macrophage.DEG_FILE; pull raw pvalue +
# log2FoldChange). DE_M0high_rd.xlsx is intentionally excluded (no counterpart elsewhere).
def load_rna(proteins=PROTEINS_TO_CALC) -> pl.DataFrame:
    # Match gene symbols case-insensitively; emit the canonical requested name so the
    # RNA rows facet together with the protein rows (e.g. C12ORF4 -> C12orf4).
    canonical = {p.upper(): p for p in proteins}
    rows = []
    for condition in CONDITIONS:
        wb = _wb(RNA_FC_P_VALUES_FOLDER / DEG_FILE[condition])
        for i, row in enumerate(wb[DEG_SHEET].iter_rows(values_only=True)):
            if i == 0:
                hdr = list(row)
                i_sym, i_pval, i_lfc = (hdr.index(c) for c in
                                        ("gene.symbol", "pvalue", "log2FoldChange"))
            else:
                sym = str(row[i_sym]).upper() if row[i_sym] is not None else None
                if sym in canonical:
                    rows.append(_row(canonical[sym], condition, "rna",
                                     row[i_pval], row[i_lfc]))
        wb.close()
    return pl.DataFrame(rows, schema=SCHEMA)


# Reactivity — residue-level. There is no p-value; a "change" is the categorical call in
# `direction_of_reactivity_change` (Higher/Lower vs unchanged). We keep only called changes
# so the plot draws a dot only where a residue actually changed. Fold change is per residue:
# log2(rc_percent_control / 100) (100% control = unchanged). p_value is left null.
def load_reactivity(proteins=PROTEINS_TO_CALC) -> pl.DataFrame:
    canonical = {p.upper(): p for p in proteins}
    return (
        pl.read_csv(REACTIVITY_FC_P_VALUES, infer_schema_length=5000,
                    ignore_errors=True)
        .filter(
            pl.col("protein").str.to_uppercase().is_in(list(canonical))
            & pl.col("condition").is_in(SHARED_CONDITIONS)
            & pl.col("direction_of_reactivity_change").is_in(["Higher", "Lower"])
        )
        .select(
            pl.col("protein").str.to_uppercase().replace(canonical).alias("protein"),
            "condition",
            # Encode the changed cysteine in the modality label, e.g. "reactivity (C737)",
            # so the heatmap can show which residue changed.
            pl.concat_str([pl.lit("reactivity ("), pl.col("residue"), pl.lit(")")])
                .alias("modality"),
            pl.lit(None, dtype=pl.Float64).alias("p_value"),
            (pl.col("rc_percent_control").cast(pl.Float64) / 100.0)
                .log(2).alias("log2_fold_change"),
        )
        .select(SCHEMA)
    )


# Western blot — one sheet per protein. Normalized rows set M0=1, so each value is a
# linear fold change vs M0; take log2, one-sample t-test vs 0. The 2026 compiled source
# file lays each replicate out as an "M0 … STING" block; blocks stack vertically (D1/D2/
# D3) and some proteins also carry rerun blocks to the right of the originals, tagged in
# the header cell just left of that block's "M0" (e.g. "RERUN - replace OG D3", "rerun 2
# (along with STAT4)", "R2"). WB_CONFIG picks, per replicate, which block to trust.
#   target: the value row's label suffix ("vinc/m0 norm", or exactly "double norm").
#   reps:   replicate -> block tag; None = leftmost/original block, else a substring of
#           the block's header note (matched case-insensitively).
WB_CONFIG = {
    "CAMK1":  {"target": "vinc/m0 norm", "reps": {"D1": None,      "D2": None, "D3": "rerun"}},
    "SCO2":   {"target": "vinc/m0 norm", "reps": {"D1": "rerun 2", "D2": None, "D3": None}},
    "STAT4":  {"target": "vinc/m0 norm", "reps": {"D1": None,      "D2": None, "D3": None}},
    "S100A4": {"target": "double norm",  "reps": {"D1": None,      "D2": None, "D3": None}},
    "CLIC1":  {"target": "vinc/m0 norm", "reps": {"D1": "R2",      "D2": None, "D3": None}},
}


def _wb_target_match(label, target) -> bool:
    if label is None:
        return False
    text = str(label).strip().lower().replace(" ", "")
    return text == "doublenorm" if target == "double norm" \
        else text.endswith("vinc/m0norm")


def _wb_candidates(sheet, target):
    """Scan a protein sheet and yield every (replicate, tag, {cond: value}, row_index)
    value row. An "M0 … STING" run starting at column c is a block; its header note sits
    at row[c-1]. Blocks re-detect on each header row so vertically stacked D1/D2/D3 and
    horizontally repeated rerun blocks are all captured."""
    active = {}  # col c -> (conditions list, tag, marker)
    for r_idx, row in enumerate(sheet.iter_rows(values_only=True)):
        m0_cols = [c for c in range(1, len(row)) if isinstance(row[c], str)
                   and row[c].strip() == "M0"]
        if m0_cols:  # header row: reset the active blocks to those found here
            active = {}
            for c in m0_cols:
                raw = "" if row[c - 1] is None else str(row[c - 1]).strip().lower()
                marker = re.search(r"d\d", raw)
                marker = marker.group(0) if marker else None
                tag = None if raw == "" or re.fullmatch(r"d\d", raw) else raw
                active[c] = ([_norm_cond(x) for x in row[c:c + 8]], tag, marker)
            continue
        for c, (conditions, tag, marker) in active.items():
            label = row[c - 1] if c - 1 < len(row) else None
            if not _wb_target_match(label, target):
                continue
            rep = marker
            if rep is None:
                m = re.search(r"d\d", str(label).lower())
                rep = m.group(0) if m else None
            values = {cond: val for cond, val in zip(conditions, row[c:c + 8])}
            yield rep, tag, values, r_idx


def load_western_blot() -> pl.DataFrame:
    wb = _wb(WESTERN_BLOT_REPLICATE_VALUES)
    rows = []
    for protein in PROTEINS_TO_CALC:
        cfg = WB_CONFIG[protein]
        candidates = list(_wb_candidates(wb[protein], cfg["target"]))
        per_cond = {}
        for slot, want_tag in cfg["reps"].items():  # slot: "D1"/"D2"/"D3"
            matches = [
                cand for cand in candidates
                if cand[0] == slot.lower()
                and (want_tag is None and cand[1] is None
                     or want_tag is not None and cand[1] is not None
                     and want_tag.lower() in cand[1])
            ]
            if not matches:
                continue
            _, _, values, _ = min(matches, key=lambda cand: cand[3])  # topmost
            for cond, val in values.items():
                if cond in SHARED_CONDITIONS and val is not None and val > 0:
                    per_cond.setdefault(cond, []).append(math.log2(val))
        for cond in SHARED_CONDITIONS:
            vals = per_cond.get(cond, [])
            p = stats.ttest_1samp(vals, 0.0).pvalue if len(vals) >= 2 else None
            fc = sum(vals) / len(vals) if vals else None
            rows.append(_row(protein, cond, "western_blot", p, fc))
    wb.close()
    return pl.DataFrame(rows, schema=SCHEMA)


# qPCR — single sheet, blocks laid out left-to-right per protein. Values are log2;
# M0 is the matched per-replicate baseline. Paired t-test of condition vs M0.
def load_qpcr() -> pl.DataFrame:
    wb = _wb(QPCR_REPLICATE_VALUES)
    grid = [list(r) for r in wb["qPCR"].iter_rows(values_only=True)]
    wb.close()
    # Row 0 has protein names above each block's first ("D1") column; row 1 has the
    # D1/D2/D3 markers. The condition labels sit in the column just left of D1.
    name_row, rep_row = grid[0], grid[1]
    rows = []
    for col, name in enumerate(name_row):
        if name not in PROTEINS_TO_CALC:
            continue
        rep_cols = [c for c in range(col, len(rep_row))
                    if isinstance(rep_row[c], str)
                    and rep_row[c].strip() in ("D1", "D2", "D3")][:3]
        label_col = rep_cols[0] - 1
        cond_vals = {}
        for r in range(2, len(grid)):
            label = grid[r][label_col] if label_col < len(grid[r]) else None
            vals = [grid[r][c] for c in rep_cols]
            if label is not None and not any(v is None for v in vals):
                cond_vals[_norm_cond(label)] = [float(v) for v in vals]
        base = cond_vals.get("M0")
        for cond in SHARED_CONDITIONS:
            vals = cond_vals.get(cond)
            if vals is None or base is None or len(vals) != len(base):
                rows.append(_row(name, cond, "qpcr", None, None))
            else:
                p = stats.ttest_rel(vals, base).pvalue
                fc = sum(v - b for v, b in zip(vals, base)) / len(vals)
                rows.append(_row(name, cond, "qpcr", p, fc))
    return pl.DataFrame(rows, schema=SCHEMA)


def main() -> None:
    parts = [load_protein(), load_rna(), load_western_blot(), load_qpcr()]
    combined = pl.concat(parts).sort(["protein", "condition", "modality"])
    combined.write_csv(OUTPUT_CSV)
    # Keep the figure-panel copy the R heatmap reads in sync with the local copy.
    MM_HEATMAP_DIR.mkdir(parents=True, exist_ok=True)
    combined.write_csv(MM_HEATMAP_DIR / "multi_modal_significance.csv")
    print(f"Wrote {combined.height} rows to {OUTPUT_CSV}")
    print(f"  and to {MM_HEATMAP_DIR / 'multi_modal_significance.csv'}")
    print(combined)

    # TBCK complex — RNA-seq + whole proteome + residue-level reactivity (no WB/qPCR).
    tbck = pl.concat([
        load_protein(TBCK_COMPLEX_PROTEINS),
        load_rna(TBCK_COMPLEX_PROTEINS),
        load_reactivity(TBCK_COMPLEX_PROTEINS),
    ]).sort(["protein", "condition", "modality"])
    TBCK_OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    tbck.write_csv(TBCK_OUTPUT_CSV)
    print(f"Wrote {tbck.height} rows to {TBCK_OUTPUT_CSV}")
    print(tbck)


if __name__ == "__main__":
    main()
